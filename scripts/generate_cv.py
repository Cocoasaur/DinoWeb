from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src" / "assets" / "resume" / "John_Arquesola_Curriculum_Vitae.docx"
BLACK = "000000"


def keep(paragraph, with_next=False):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepLines")) is None:
        p_pr.append(OxmlElement("w:keepLines"))
    if with_next and p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def section_heading(doc, text):
    paragraph = doc.add_paragraph(style="Section")
    paragraph.add_run(text.upper())
    keep(paragraph, with_next=True)


def dated_line(doc, left, right, style="Entry"):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    left_run = paragraph.add_run(left)
    left_run.bold = True
    paragraph.add_run("\t")
    right_run = paragraph.add_run(right)
    right_run.bold = True
    keep(paragraph, with_next=True)
    return paragraph


def bullet(doc, text):
    paragraph = doc.add_paragraph("•  " + text, style="Bullet")
    keep(paragraph)


def skill(doc, label, text):
    paragraph = doc.add_paragraph(style="Body")
    run = paragraph.add_run(label + ": ")
    run.bold = True
    paragraph.add_run(text)
    keep(paragraph)


def project(doc, name, role, url, link_text, bullets):
    paragraph = doc.add_paragraph(style="Entry")
    name_run = paragraph.add_run(name)
    name_run.bold = True
    paragraph.add_run(" | " + role)
    if url:
        paragraph.add_run(" | ")
        hyperlink(paragraph, link_text, url)
    keep(paragraph, with_next=True)
    for text in bullets:
        bullet(doc, text)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.15)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    body = doc.styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)
    body.font.name = "Times New Roman"
    body.font.size = Pt(9.4)
    body.font.color.rgb = RGBColor.from_string(BLACK)
    body.paragraph_format.space_after = Pt(1)
    body.paragraph_format.line_spacing = 1.02

    entry = doc.styles.add_style("Entry", WD_STYLE_TYPE.PARAGRAPH)
    entry.font.name = "Times New Roman"
    entry.font.size = Pt(9.5)
    entry.font.color.rgb = RGBColor.from_string(BLACK)
    entry.paragraph_format.space_after = Pt(0.8)
    entry.paragraph_format.line_spacing = 1.02

    bullet_style = doc.styles.add_style("Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet_style.font.name = "Times New Roman"
    bullet_style.font.size = Pt(9.15)
    bullet_style.font.color.rgb = RGBColor.from_string(BLACK)
    bullet_style.paragraph_format.left_indent = Inches(0.14)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.14)
    bullet_style.paragraph_format.space_after = Pt(0.8)
    bullet_style.paragraph_format.line_spacing = 1.02

    heading = doc.styles.add_style("Section", WD_STYLE_TYPE.PARAGRAPH)
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(10.2)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(BLACK)
    heading.paragraph_format.space_before = Pt(5)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.line_spacing = 1.0
    p_pr = heading.element.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLACK)
    borders.append(bottom)
    p_pr.append(borders)


def build_cv():
    doc = Document()
    configure(doc)
    doc.sections[0].top_margin = Inches(0.6)
    doc.sections[0].bottom_margin = Inches(0.6)
    doc.styles["Normal"].font.size = Pt(9.6)
    doc.styles["Body"].font.size = Pt(9.6)
    doc.styles["Body"].paragraph_format.space_after = Pt(3.2)
    doc.styles["Body"].paragraph_format.line_spacing = 1.08
    doc.styles["Entry"].font.size = Pt(9.8)
    doc.styles["Entry"].paragraph_format.space_after = Pt(3.2)
    doc.styles["Entry"].paragraph_format.line_spacing = 1.08
    doc.styles["Bullet"].font.size = Pt(9.45)
    doc.styles["Bullet"].paragraph_format.space_after = Pt(3.3)
    doc.styles["Bullet"].paragraph_format.line_spacing = 1.08
    doc.styles["Section"].font.size = Pt(10.5)
    doc.styles["Section"].paragraph_format.space_before = Pt(12)
    doc.styles["Section"].paragraph_format.space_after = Pt(5)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1)
    run = name.add_run("JOHN JESSIENEL M. ARQUESOLA")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(17)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(0)
    contact.add_run("Iloilo City, Philippines | +63 961 557 0380 | ")
    hyperlink(contact, "cocoasaurjl@gmail.com", "mailto:cocoasaurjl@gmail.com")

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_after = Pt(1)
    hyperlink(links, "LinkedIn", "https://linkedin.com/in/dinoweb")
    links.add_run(" | ")
    hyperlink(links, "GitHub", "https://github.com/Cocoasaur")
    links.add_run(" | ")
    hyperlink(links, "Portfolio", "https://cocoasaur.github.io/DinoWeb/")

    section_heading(doc, "Education")
    dated_line(doc, "University of San Agustin — Bachelor of Science in Computer Science", "Expected June 2027")
    p = doc.add_paragraph("Iloilo City, Philippines | Academic focus: software development, algorithms, and data structures", style="Body")
    keep(p)
    dated_line(doc, "St. Anthony’s College — Senior High School, STEM", "2021–2023")
    dated_line(doc, "Dao Catholic High School, Inc. — Junior High School and Elementary", "2011–2020")

    section_heading(doc, "Technical Skills")
    skill(doc, "Languages", "JavaScript, TypeScript, Python, Java, GDScript")
    skill(doc, "Frontend", "React, Three.js, Tailwind CSS, Tkinter, HTML5, CSS3, Vite")
    skill(doc, "Backend and Databases", "Node.js, Firebase, Firestore, MySQL, SQLite")
    skill(doc, "Tools", "Git, GitHub, VS Code, Postman, Figma, Draw.io, Godot Engine")

    section_heading(doc, "Leadership")
    dated_line(doc, "Augustinian Developer Society — Quality Assurance Lead", "Aug 2025–Present")
    bullet(doc, "Served as Technical Judge for the DevDash Hackathon and audited competing teams’ GitHub repositories for compliance.")
    bullet(doc, "Established development standards and reviewed code for performance and user-experience quality.")
    dated_line(doc, "Augustinian Developer Society — Junior Developer", "Oct 2024–Aug 2025")
    bullet(doc, "Supported organization events as part of the technical team, providing setup and on-site technical assistance.")
    bullet(doc, "Participated in hackathons to develop technology-based solutions to real-world problems.")

    section_heading(doc, "Projects")
    project(doc, "Quarto — Boarding House Management System", "Full-Stack Developer", "https://quarto-4d613.web.app", "Live Project", [
        "Built administrator and tenant portals for rooms, payments, tickets, announcements, approvals, and reports using JavaScript and Firebase.",
        "Implemented Cloud Functions, transactions, caching, indexed queries, role-based security rules, validation, audit logs, and soft-delete recovery.",
    ])
    project(doc, "DinoWeb — Personal Web Portfolio", "Full-Stack Developer", "https://cocoasaur.github.io/DinoWeb/", "Portfolio", [
        "Built and deployed a responsive React and Three.js portfolio featuring real-time 3D rendering, animation, and performance-aware loading.",
    ])
    project(doc, "Earth Arcade", "Full-Stack Developer", "https://github.com/FlimsyOwl12/GAME-ON-Hackathon-Earth-Arcade.git", "GitHub", [
        "Developed UI/UX, scene transitions, 2D environments, and gameplay mechanics for a Godot game promoting climate action under UN SDG 13.",
    ])
    project(doc, "TIP Airlines", "Full-Stack Developer", None, None, [
        "Created a Python, Tkinter, and SQLite airline system with passenger booking, flight scheduling, and persistent data storage.",
    ])

    section_heading(doc, "Certifications")
    bullet(doc, "Introduction to Data Science — Cisco Networking Academy, Oct 2025")
    bullet(doc, "Microsoft Excel guided projects — Coursera Project Network, Jun 2025")
    bullet(doc, "Code a Joke-Telling Talkbot — Google for Education, Jun 2025 | A.I.GNITE — NVIDIA AI Academy Philippines, Mar 2025")
    bullet(doc, "Basic Web Development Workshop — Zuitt, Nov 2023 | Google DevFest — GDG Bacolod, Oct 2023")

    section_heading(doc, "Activities")
    dated_line(doc, "National AI Hackathon — Finalist", "Aug 2025")
    dated_line(doc, "DEVCON Iloilo — Volunteer Publicity-Material Designer", "Oct 2024")

    doc.core_properties.title = "Curriculum Vitae — John Jessienel M. Arquesola"
    doc.core_properties.subject = "One-page ATS-friendly curriculum vitae"
    doc.core_properties.author = "John Jessienel M. Arquesola"
    doc.core_properties.keywords = "curriculum vitae, computer science, software development"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_cv()
