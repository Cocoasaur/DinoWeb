from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src" / "assets" / "resume" / "John_Arquesola_Resume.docx"
BLACK = "000000"
LINK_COLOR = BLACK


def keep(paragraph, with_next=False):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepLines")) is None:
        p_pr.append(OxmlElement("w:keepLines"))
    if with_next and p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def add_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLACK)
    borders.append(bottom)
    p_pr.append(borders)


def section_heading(doc, text):
    p = doc.add_paragraph(style="Section")
    p.add_run(text.upper())
    add_rule(p)
    keep(p, with_next=True)


def aligned_line(doc, left, right="", style="Entry", bold=False, italic=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.05), WD_TAB_ALIGNMENT.RIGHT)
    left_run = p.add_run(left)
    left_run.bold = bold
    left_run.italic = italic
    if right:
        p.add_run("\t")
        right_run = p.add_run(right)
        right_run.bold = bold
        right_run.italic = italic
    keep(p, with_next=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph("•  " + text, style="Bullet")
    keep(p)


def skill(doc, label, text):
    p = doc.add_paragraph(style="Body")
    lead = p.add_run(label + ": ")
    lead.bold = True
    p.add_run(text)
    keep(p)


def leadership(doc, organization, role, dates, bullets):
    aligned_line(doc, organization, "Iloilo City, Philippines", bold=True)
    aligned_line(doc, role, dates, italic=True)
    for text in bullets:
        bullet(doc, text)


def project(doc, name, role, bullets, url=None, label=None):
    p = doc.add_paragraph(style="Entry")
    title = p.add_run(name)
    title.bold = True
    if url:
        p.add_run(" | ")
        hyperlink(p, label or url, url)
    keep(p, with_next=True)
    role_line = doc.add_paragraph(role, style="Meta")
    role_line.runs[0].italic = True
    keep(role_line, with_next=True)
    for text in bullets:
        bullet(doc, text)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    body = doc.styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)
    body.font.name = "Times New Roman"
    body.font.size = Pt(10.2)
    body.paragraph_format.space_after = Pt(2)
    body.paragraph_format.line_spacing = 1.08

    entry = doc.styles.add_style("Entry", WD_STYLE_TYPE.PARAGRAPH)
    entry.font.name = "Times New Roman"
    entry.font.size = Pt(10.3)
    entry.paragraph_format.space_after = Pt(1.5)
    entry.paragraph_format.line_spacing = 1.08

    meta = doc.styles.add_style("Meta", WD_STYLE_TYPE.PARAGRAPH)
    meta.font.name = "Times New Roman"
    meta.font.size = Pt(9.9)
    meta.paragraph_format.space_after = Pt(1.2)
    meta.paragraph_format.line_spacing = 1.06

    bullet_style = doc.styles.add_style("Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet_style.font.name = "Times New Roman"
    bullet_style.font.size = Pt(9.65)
    bullet_style.paragraph_format.left_indent = Inches(0.16)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.11)
    bullet_style.paragraph_format.space_after = Pt(2)
    bullet_style.paragraph_format.line_spacing = 1.08

    section_style = doc.styles.add_style("Section", WD_STYLE_TYPE.PARAGRAPH)
    section_style.font.name = "Times New Roman"
    section_style.font.size = Pt(11)
    section_style.font.bold = True
    section_style.paragraph_format.space_before = Pt(9)
    section_style.paragraph_format.space_after = Pt(3.5)
    section_style.paragraph_format.line_spacing = 1.06


def build_resume():
    doc = Document()
    configure(doc)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1.2)
    run = name.add_run("John Jessienel M. Arquesola")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(0.8)
    contact.add_run("Iloilo City, Philippines  ·  ")
    hyperlink(contact, "cocoasaurjl@gmail.com", "mailto:cocoasaurjl@gmail.com")
    contact.add_run("  ·  +63 961 557 0380")

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_after = Pt(3)
    hyperlink(links, "LinkedIn", "https://linkedin.com/in/dinoweb")
    links.add_run("  ·  ")
    hyperlink(links, "GitHub", "https://github.com/Cocoasaur")
    links.add_run("  ·  ")
    hyperlink(links, "Portfolio", "https://cocoasaur.github.io/DinoWeb/")

    section_heading(doc, "Education")
    aligned_line(doc, "University of San Agustin", "Iloilo City, Philippines", bold=True)
    aligned_line(doc, "Bachelor of Science in Computer Science", "Expected June 2027", italic=True)
    aligned_line(doc, "St. Anthony’s College — Senior High School, STEM", "2021–2023", bold=True)
    aligned_line(doc, "Dao Catholic High School, Inc. — Junior High School and Elementary", "2011–2020", bold=True)

    section_heading(doc, "Technical Skills")
    skill(doc, "Languages", "JavaScript, TypeScript, Python, Java, GDScript")
    skill(doc, "Frontend", "React, Three.js, Tailwind CSS, Tkinter, HTML5, CSS3, Vite")
    skill(doc, "Backend and Databases", "Node.js, Firebase, Firestore, MySQL, SQLite")
    skill(doc, "Tools", "Git, GitHub, VS Code, Postman, Figma, Draw.io, Godot Engine")

    section_heading(doc, "Leadership")
    leadership(doc, "Augustinian Developer Society — University of San Agustin", "Quality Assurance Lead", "Aug 2025–Present", [
        "Served as Technical Judge for the DevDash Hackathon and audited competing teams’ GitHub repositories for compliance.",
        "Established development standards and reviewed code for performance and user-experience quality.",
    ])
    leadership(doc, "Augustinian Developer Society — University of San Agustin", "Junior Developer", "Oct 2024–Aug 2025", [
        "Supported organization events as part of the technical team, providing setup and on-site technical assistance.",
        "Participated in hackathons to develop technology-based solutions to real-world problems.",
    ])

    section_heading(doc, "Projects")
    project(doc, "DinoWeb — Personal Web Portfolio", "Full-Stack Developer", [
        "Built an interactive React and Three.js portfolio featuring real-time 3D rendering and animation.",
        "Deployed and maintained the application on GitHub Pages.",
    ], "https://cocoasaur.github.io/DinoWeb/", "cocoasaur.github.io/DinoWeb")
    project(doc, "Quarto — Boarding House Management System", "Full-Stack Developer", [
        "Built administrator and tenant portals for rooms, payments, tickets, announcements, approvals, and reports.",
        "Implemented Cloud Functions, transactions, caching, indexed queries, role-based security rules, validation, and audit logs.",
    ], "https://quarto-4d613.web.app", "quarto-4d613.web.app")
    project(doc, "Earth Arcade", "Full-Stack Developer", [
        "Developed UI/UX, 2D environments, and gameplay mechanics for a Godot game promoting climate action under UN SDG 13.",
        "Coordinated repository structure and version control for the team project.",
    ], "https://github.com/FlimsyOwl12/GAME-ON-Hackathon-Earth-Arcade.git", "GitHub repository")
    project(doc, "TIP Airlines", "Full-Stack Developer", [
        "Created a Python, Tkinter, and SQLite airline system with passenger booking and flight scheduling.",
        "Applied database fundamentals for persistent storage and data integrity.",
    ])

    section_heading(doc, "Activities")
    aligned_line(doc, "National AI Hackathon", "Iloilo City, Philippines", bold=True)
    aligned_line(doc, "Finalist", "Aug 2025", italic=True)
    aligned_line(doc, "DEVCON Iloilo", "Iloilo City, Philippines", bold=True)
    aligned_line(doc, "Volunteer Publicity-Material Designer", "Oct 2024", italic=True)

    doc.core_properties.title = "Resume_Arquesola"
    doc.core_properties.subject = "Software development resume"
    doc.core_properties.author = "John Jessienel M. Arquesola"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_resume()
